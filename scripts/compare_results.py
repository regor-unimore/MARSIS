"""Compare the results in output/ with the values reported in results/Computational Results.xlsx.

For each experiment of results/experiments.csv whose output is available, the objective values (and upper
bounds) are compared with the corresponding rows of the workbook:
  - H1, H2, the Lagrangian relaxation and the ILP solved to optimality are deterministic: their values must
    match exactly (MATCH / MISMATCH);
  - the ILP runs that reached the time limit, and the matheuristics, depend on the computing time and thus on
    the hardware: their values are reported together with the relative difference (TIME-DEPENDENT).

Usage examples:
    python scripts/compare_results.py
    python scripts/compare_results.py --sheet Heuristics --report output/comparison.csv
"""
import argparse
import csv
import sys
from pathlib import Path
from statistics import mean
import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[1]

parser = argparse.ArgumentParser(description="Compare the results with the results workbook",
                                 formatter_class=argparse.ArgumentDefaultsHelpFormatter)
parser.add_argument('--manifest', type=Path, default=REPO_ROOT / "results" / "experiments.csv", help="Experiment manifest")
parser.add_argument('--workbook', type=Path, default=REPO_ROOT / "results" / "Computational Results.xlsx",
                    help="Workbook with the results of the paper")
parser.add_argument('--id', nargs='+', help="Compare only the experiments with these ids")
parser.add_argument('--sheet', nargs='+', help="Compare only the experiments of these workbook sheets")
parser.add_argument('--algorithm', nargs='+', help="Compare only these algorithms")
parser.add_argument('--instance', nargs='+', help="Compare only the experiments on these instances")
parser.add_argument('--all', action='store_true', help="Also list the experiments whose output is missing")
parser.add_argument('--report', type=Path, help="Save the comparison to this CSV file")
args = parser.parse_args()


with open(args.manifest, newline='', encoding='utf-8') as file:
    manifest = list(csv.DictReader(file))
filters = {'id': args.id, 'sheet': args.sheet, 'algorithm': args.algorithm, 'instance': args.instance}
selected = [e for e in manifest if all(v is None or e[k] in v for k, v in filters.items())]

workbook = openpyxl.load_workbook(args.workbook, data_only=True, read_only=True)
sheets = {name: [None, None] + [row for row in workbook[name].iter_rows(min_row=2, values_only=True)]
          for name in workbook.sheetnames}  # indexed by workbook row number


def workbook_rows(experiment):
    first, _, last = experiment['workbook_rows'].partition('-')
    return [sheets[experiment['sheet']][i] for i in range(int(first), int(last or first) + 1)]


def read_values(path, key):
    """Values of the summary lines '<key> = <value>' of an output file."""
    values = []
    with open(path, 'r') as file:
        for line in file:
            name, sep, value = line.partition('=')
            if sep and name.strip() == key:
                values.append(float(value))
    return values


def equal(a, b):
    return abs(a - b) <= 1e-6 * max(1.0, abs(b))


def equal_bound(a, b):
    """Upper bounds are rounded in the workbook: compare them with the same number of decimals."""
    decimals = len(repr(float(b)).split('.')[1].rstrip('0'))
    return equal(round(a, decimals), b)


results = []
for e in selected:
    output = REPO_ROOT / e['output_files'].split(';')[0]
    row = {'id': e['id'], 'sheet': e['sheet'], 'workbook_rows': e['workbook_rows'], 'algorithm': e['algorithm'],
           'instance': e['instance'], 'N': e['N'], 'variant': e['variant'], 'policy': e['policy'],
           'repetitions': e['repetitions'], 'paper': '', 'reproduced': '', 'difference': '', 'status': ''}
    wb = workbook_rows(e)
    if not output.is_file():
        row['status'] = 'MISSING'
    elif e['algorithm'] in ('H1', 'H2'):
        paper, reproduced = wb[0][6], read_values(output, 'objective function')[-1]
        row.update(paper=paper, reproduced=reproduced, status='MATCH' if equal(reproduced, paper) else 'MISMATCH')
    elif e['algorithm'] == 'Lagrangian':
        paper = (wb[0][5], wb[0][6])
        reproduced = (read_values(output, 'objective function')[-1], read_values(output, 'UB')[-1])
        row.update(paper=f"obj={paper[0]} UB={paper[1]}", reproduced=f"obj={reproduced[0]} UB={reproduced[1]}",
                   status='MATCH' if equal(reproduced[0], paper[0]) and equal_bound(reproduced[1], paper[1]) else 'MISMATCH')
    elif e['algorithm'] == 'ILP':
        lb, ub, time = wb[0][5], wb[0][6], wb[0][7]
        obj, bound = read_values(output, 'objective function')[-1], read_values(output, 'UB')[-1]
        row.update(paper=f"obj={lb} UB={ub}", reproduced=f"obj={obj} UB={bound}")
        if isinstance(time, (int, float)):  # solved to optimality
            row['status'] = 'MATCH' if equal(obj, lb) and equal_bound(bound, ub) else 'MISMATCH'
        else:  # time limit or memory limit
            row['status'] = 'TIME-DEPENDENT'
            if isinstance(lb, (int, float)) and lb > 0:
                row['difference'] = f"{100 * (obj - lb) / lb:+.3f}%"
    else:  # matheuristics: one objective value per repetition
        paper = [r[7] for r in wb]
        reproduced = read_values(output, 'Objective Value')[-len(paper):]
        row.update(paper=f"avg={mean(paper):.2f} ({len(paper)} runs)",
                   reproduced=f"avg={mean(reproduced):.2f} ({len(reproduced)} runs)",
                   difference=f"{100 * (mean(reproduced) - mean(paper)) / mean(paper):+.3f}%",
                   status='TIME-DEPENDENT' if len(reproduced) == len(paper) else 'INCOMPLETE')
    results.append(row)

shown = [r for r in results if args.all or r['status'] != 'MISSING']
for r in shown:
    config = ' '.join(x for x in (r['instance'], f"N={r['N']}", r['variant'], f"P{r['policy']}" if r['policy'] else '',
                                  r['repetitions']) if x)
    print(f"{r['id']} {r['algorithm']:<12} {config:<40} paper: {r['paper']!s:<32} reproduced: {r['reproduced']!s:<32} "
          f"{r['difference']:>9} {r['status']}")

counts = {s: sum(r['status'] == s for r in results) for s in ('MATCH', 'MISMATCH', 'TIME-DEPENDENT', 'INCOMPLETE', 'MISSING')}
print('\n' + ', '.join(f"{k}: {v}" for k, v in counts.items()))

if args.report and results:
    args.report.parent.mkdir(parents=True, exist_ok=True)
    with open(args.report, 'w', newline='', encoding='utf-8') as file:
        writer = csv.DictWriter(file, fieldnames=list(results[0].keys()))
        writer.writeheader()
        writer.writerows(results)
    print(f"Comparison saved to {args.report}")

sys.exit(1 if counts['MISMATCH'] else 0)
